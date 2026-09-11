from __future__ import annotations

import os
from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.crud.report import (
    _MOCK_LEDGER,
    _MOCK_SUMMARY,
    get_financial_summary,
    get_order_ledger,
)
from app.db.session import get_db
from app.deps import get_current_staff, require_role
from app.models.staff import StaffRole
from app.schemas.report import FinancialSummary, OrderLedgerRow

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    _HAS_REPORTLAB = True
except ImportError:
    _HAS_REPORTLAB = False

router = APIRouter(prefix="/reports", tags=["reports"])

# Controle de mock via env — útil em CI e desenvolvimento sem banco configurado
_USE_MOCK = os.getenv("REPORTS_USE_MOCK", "false").lower() == "true"


# ─── Helpers de dados ─────────────────────────────────────────────────────────

def _get_summary(db: Session) -> FinancialSummary:
    if _USE_MOCK:
        return _MOCK_SUMMARY
    try:
        return get_financial_summary(db)
    except Exception:
        return _MOCK_SUMMARY


def _get_ledger(db: Session) -> list[OrderLedgerRow]:
    if _USE_MOCK:
        return _MOCK_LEDGER
    try:
        return get_order_ledger(db)
    except Exception:
        return _MOCK_LEDGER


# ─── Endpoints JSON ───────────────────────────────────────────────────────────

@router.get("/summary", response_model=FinancialSummary)
def financial_summary(
    db: Session = Depends(get_db),
    _staff=Depends(require_role([StaffRole.admin])),
) -> FinancialSummary:
    """Resumo financeiro consolidado para os cards de KPI e gráficos."""
    return _get_summary(db)


@router.get("/ledger", response_model=list[OrderLedgerRow])
def order_ledger(
    db: Session = Depends(get_db),
    _staff=Depends(require_role([StaffRole.admin])),
) -> list[OrderLedgerRow]:
    """Lista linha-a-linha dos pedidos entregues com custo e lucro."""
    return _get_ledger(db)


# ─── Export Excel ─────────────────────────────────────────────────────────────

@router.get("/export.xlsx")
def export_xlsx(
    db: Session = Depends(get_db),
    _staff=Depends(require_role([StaffRole.admin])),
) -> StreamingResponse:
    """
    Gera e retorna uma planilha Excel em memória com duas abas:
      1. Resumo Financeiro — KPIs e breakdown por canal/staff
      2. Lista de Pedidos  — ledger linha a linha

    Requer openpyxl instalado (adicionar a requirements.txt).
    """
    if not _HAS_OPENPYXL:
        from fastapi import HTTPException
        raise HTTPException(
            status_code=503,
            detail="openpyxl não está instalado no servidor. Adicione ao requirements.txt.",
        )

    summary = _get_summary(db)
    ledger  = _get_ledger(db)
    buffer  = _build_workbook(summary, ledger)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="relatorio_pizzashop.xlsx"'},
    )


# ─── Export PDF ───────────────────────────────────────────────────────────────

@router.get("/export.pdf")
def export_pdf(
    db: Session = Depends(get_db),
    _staff=Depends(require_role([StaffRole.admin])),
) -> StreamingResponse:
    """
    Gera e retorna um PDF em memória com o mesmo conteúdo do Excel:
    resumo financeiro (KPIs, por canal, por staff) + lista de pedidos.
    Requer reportlab instalado (adicionar a requirements.txt).
    """
    if not _HAS_REPORTLAB:
        raise HTTPException(
            status_code=503,
            detail="reportlab não está instalado no servidor. Adicione ao requirements.txt.",
        )

    summary = _get_summary(db)
    ledger = _get_ledger(db)
    buffer = _build_pdf(summary, ledger)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="relatorio_pizzashop.pdf"'},
    )


# ─── Builder da planilha ─────────────────────────────────────────────────────

def _build_workbook(summary: FinancialSummary, ledger: list[OrderLedgerRow]) -> BytesIO:
    """Cria o workbook com estilização básica e retorna um BytesIO pronto para streaming."""
    wb = openpyxl.Workbook()

    _build_sheet_summary(wb, summary)
    _build_sheet_ledger(wb, ledger)

    # Remove a sheet padrão vazia criada pelo openpyxl
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# Cores
_HEADER_FILL   = PatternFill("solid", fgColor="C0392B")   # vermelho pizzaria
_SECTION_FILL  = PatternFill("solid", fgColor="F5B7B1")   # rosa claro
_HEADER_FONT   = Font(bold=True, color="FFFFFF", size=11)
_SECTION_FONT  = Font(bold=True, color="922B21", size=10)
_BOLD_FONT     = Font(bold=True)


def _style_header_row(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _currency(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _build_sheet_summary(wb, summary: FinancialSummary) -> None:
    ws = wb.create_sheet("Resumo Financeiro")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    def write(row, label, value, bold=False):
        ws.cell(row=row, column=1, value=label).font = _BOLD_FONT if bold else Font()
        ws.cell(row=row, column=2, value=value)

    # ── KPIs ──
    ws.cell(row=1, column=1, value="RESUMO FINANCEIRO — PizzaShop")
    ws.cell(row=1, column=1).font = Font(bold=True, size=13, color="C0392B")
    ws.merge_cells("A1:B1")

    ws.cell(row=2, column=1, value="Indicadores Gerais").fill = _SECTION_FILL
    ws.cell(row=2, column=1).font = _SECTION_FONT
    ws.merge_cells("A2:B2")

    write(3,  "Faturamento Total",   _currency(summary.total_revenue),  bold=True)
    write(4,  "Custo Total",         _currency(summary.total_cost))
    write(5,  "Lucro Bruto",         _currency(summary.gross_profit),   bold=True)
    write(6,  "Total de Pedidos",    summary.orders_count)
    rating_str = f"{summary.average_rating:.1f} ★" if summary.average_rating else "N/A"
    write(7,  "Avaliação Média",     rating_str)

    # ── Canal ──
    ws.cell(row=9, column=1, value="Por Canal de Venda").fill = _SECTION_FILL
    ws.cell(row=9, column=1).font = _SECTION_FONT
    ws.merge_cells("A9:B9")

    headers = ["Canal", "Faturamento", "Custo", "Lucro"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=10, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    ws.column_dimensions[get_column_letter(3)].width = 16
    ws.column_dimensions[get_column_letter(4)].width = 16

    for i, ch in enumerate(summary.channel_breakdown, start=11):
        label = "Delivery" if ch.channel == "delivery" else "Presencial"
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=_currency(ch.revenue))
        ws.cell(row=i, column=3, value=_currency(ch.cost))
        ws.cell(row=i, column=4, value=_currency(ch.profit))

    # ── Staff ──
    base = 11 + len(summary.channel_breakdown) + 2
    ws.cell(row=base, column=1, value="Por Entregador / Garçom").fill = _SECTION_FILL
    ws.cell(row=base, column=1).font = _SECTION_FONT
    ws.merge_cells(f"A{base}:B{base}")

    s_headers = ["Nome", "Função", "Pedidos", "Faturamento"]
    for c, h in enumerate(s_headers, 1):
        cell = ws.cell(row=base + 1, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for i, st in enumerate(summary.staff_breakdown, start=base + 2):
        role_label = "Entregador" if st.role == "entrega" else "Garçom"
        ws.cell(row=i, column=1, value=st.name)
        ws.cell(row=i, column=2, value=role_label)
        ws.cell(row=i, column=3, value=st.orders_count)
        ws.cell(row=i, column=4, value=_currency(st.revenue))


def _build_sheet_ledger(wb, ledger: list[OrderLedgerRow]) -> None:
    ws = wb.create_sheet("Lista de Pedidos")

    headers = [
        "ID Pedido", "Data/Hora", "Canal", "Cliente",
        "Cozinheiro", "Entregador",
        "Subtotal", "Taxa Entrega", "Total", "Custo", "Lucro",
        "Avaliação", "Status",
    ]
    col_widths = [12, 20, 12, 22, 16, 16, 14, 14, 14, 14, 14, 10, 18]

    for c, (h, w) in enumerate(zip(headers, col_widths), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(ledger, start=2):
        channel_label = "Delivery" if row.channel == "delivery" else "Presencial"
        values = [
            row.order_id,
            row.created_at.replace("T", " ").replace("Z", ""),
            channel_label,
            row.customer_name,
            row.cook_name or "—",
            row.driver_name or "—",
            _currency(row.subtotal),
            _currency(row.delivery_fee),
            _currency(row.total),
            _currency(row.cost),
            _currency(row.profit),
            f"{row.rating:.1f} ★" if row.rating is not None else "—",
            row.status,
        ]
        for c, v in enumerate(values, 1):
            ws.cell(row=r, column=c, value=v)


# ─── Builder do PDF ───────────────────────────────────────────────────────────

_PDF_RED = colors.HexColor("#C0392B") if _HAS_REPORTLAB else None
_PDF_RED_DARK = colors.HexColor("#922B21") if _HAS_REPORTLAB else None
_PDF_STRIPE = colors.HexColor("#FBEEEE") if _HAS_REPORTLAB else None
_PDF_GRID = colors.HexColor("#E0E0E0") if _HAS_REPORTLAB else None


def _pdf_table_style(font_size: int = 8, header: bool = True) -> "TableStyle":
    style = [
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("GRID", (0, 0), (-1, -1), 0.25, _PDF_GRID),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        style += [
            ("BACKGROUND", (0, 0), (-1, 0), _PDF_RED),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, _PDF_STRIPE]),
        ]
    return TableStyle(style)


def _build_pdf(summary: FinancialSummary, ledger: list[OrderLedgerRow]) -> BytesIO:
    """Monta o PDF (resumo + lista de pedidos) e retorna um BytesIO pronto pra streaming."""
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.3 * cm, bottomMargin=1.3 * cm,
        title="Relatório Financeiro — PizzaShop",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("PizzaTitle", parent=styles["Title"], textColor=_PDF_RED, fontSize=18)
    section_style = ParagraphStyle(
        "PizzaSection", parent=styles["Heading2"], textColor=_PDF_RED_DARK,
        fontSize=12, spaceBefore=14, spaceAfter=6,
    )
    meta_style = ParagraphStyle("PizzaMeta", parent=styles["Normal"], textColor=colors.HexColor("#666666"))

    elements = [
        Paragraph("Relatório Financeiro — PizzaShop", title_style),
        Paragraph(
            datetime.now(timezone.utc).strftime("Gerado em %d/%m/%Y às %H:%M UTC"),
            meta_style,
        ),
        Spacer(1, 10),
    ]

    # ── Indicadores gerais ──
    elements.append(Paragraph("Indicadores gerais", section_style))
    rating_str = f"{summary.average_rating:.1f} ★" if summary.average_rating is not None else "N/A"
    kpi_rows = [
        ["Faturamento total", _currency(summary.total_revenue)],
        ["Custo total", _currency(summary.total_cost)],
        ["Lucro bruto", _currency(summary.gross_profit)],
        ["Total de pedidos", str(summary.orders_count)],
        ["Avaliação média", rating_str],
    ]
    kpi_table = Table(kpi_rows, colWidths=[6 * cm, 6 * cm])
    kpi_table.setStyle(_pdf_table_style(header=False))
    elements.append(kpi_table)

    # ── Por canal ──
    elements.append(Paragraph("Por canal de venda", section_style))
    channel_rows = [["Canal", "Faturamento", "Custo", "Lucro"]] + [
        [
            "Delivery" if ch.channel == "delivery" else "Presencial",
            _currency(ch.revenue), _currency(ch.cost), _currency(ch.profit),
        ]
        for ch in summary.channel_breakdown
    ] or [["Canal", "Faturamento", "Custo", "Lucro"]]
    channel_table = Table(channel_rows, colWidths=[5 * cm, 4.5 * cm, 4.5 * cm, 4.5 * cm])
    channel_table.setStyle(_pdf_table_style())
    elements.append(channel_table)

    # ── Por staff ──
    elements.append(Paragraph("Por entregador / garçom", section_style))
    staff_rows = [["Nome", "Função", "Pedidos", "Faturamento"]] + [
        [
            st.name, "Entregador" if st.role == "entrega" else "Garçom",
            str(st.orders_count), _currency(st.revenue),
        ]
        for st in summary.staff_breakdown
    ]
    staff_table = Table(staff_rows, colWidths=[5 * cm, 4.5 * cm, 3 * cm, 4.5 * cm])
    staff_table.setStyle(_pdf_table_style())
    elements.append(staff_table)

    # ── Lista de pedidos ──
    elements.append(Paragraph("Lista de pedidos", section_style))
    ledger_header = [
        "Pedido", "Data/Hora", "Canal", "Cliente",
        "Cozinheiro", "Entregador", "Total", "Custo", "Lucro", "Nota", "Status",
    ]
    ledger_rows = [ledger_header]
    for row in ledger:
        ledger_rows.append([
            row.order_id[:8],
            row.created_at.replace("T", " ").replace("Z", "")[:16],
            "Delivery" if row.channel == "delivery" else "Presencial",
            row.customer_name,
            row.cook_name or "—",
            row.driver_name or "—",
            _currency(row.total),
            _currency(row.cost),
            _currency(row.profit),
            f"{row.rating:.1f}★" if row.rating is not None else "—",
            row.status,
        ])
    ledger_table = Table(
        ledger_rows,
        colWidths=[2 * cm, 3 * cm, 2.3 * cm, 4 * cm, 3 * cm, 3 * cm, 2.3 * cm, 2.3 * cm, 2.3 * cm, 1.8 * cm, 2.8 * cm],
        repeatRows=1,
    )
    ledger_table.setStyle(_pdf_table_style(font_size=7))
    elements.append(ledger_table)

    doc.build(elements)
    buf.seek(0)
    return buf
